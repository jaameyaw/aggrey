#!/usr/bin/env python3
"""Generate db/seed.sql from data/Aggrey_MVP_Data.xlsx (roadmap M2).

The Excel workbook is the single source of truth for Aggrey's dining data. This
script turns it into db/seed.sql deterministically: the same workbook always
produces byte-identical SQL. That makes future menu changes a reviewable diff
instead of a hand-edit.

The workbook is validated before anything is written (see validate()); the generator
refuses to produce a seed from invalid data.

Usage (needs openpyxl):
  python3 db/generate_seed.py            # write db/seed.sql from the workbook
  python3 db/generate_seed.py --check    # fail (exit 1) if the committed db/seed.sql
                                         # differs from the workbook -- for CI
  python3 db/generate_seed.py --check-mirror
                                         # fail (exit 1) if data/Aggrey_MVP_Data.md
                                         # differs from the workbook -- for CI
"""

from __future__ import annotations

import argparse
import difflib
import sys
from collections import Counter
from datetime import date
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = REPO_ROOT / "data" / "Aggrey_MVP_Data.xlsx"
SEED_FILE = REPO_ROOT / "db" / "seed.sql"
MIRROR_FILE = REPO_ROOT / "data" / "Aggrey_MVP_Data.md"

# Markdown mirror section heading -> workbook sheet name.
SECTION_TO_SHEET = {
    "Menus": "menus",
    "Hours": "hours",
    "Holidays and breaks": "holidays",
    "Metadata": "metadata",
}

# Menus are written in batches of this many rows, matching the committed seed.
MENU_BATCH_SIZE = 85

# Expected column headers per sheet (order matters for reading).
COLUMNS = {
    "menus": ["menu_scope", "day_of_week", "meal_type", "station", "food_item", "display_order"],
    "hours": ["schedule_type", "meal_type", "start_time", "end_time"],
    "holidays": ["name", "start_date", "end_date", "schedule_type", "verification_status", "notes"],
    "metadata": ["key", "value", "notes"],
}

# Allowed enum values, mirroring db/schema.sql.
MEAL_TYPES = {"breakfast", "lunch", "lite_lunch", "dinner", "brunch"}
SCHEDULE_TYPES = {"weekday", "weekend", "holiday", "break"}
MENU_SCOPES = {"weekly", "holiday_sample"}

# This generator is pinned to one dataset version. If the data changes on
# purpose, bump DATASET_VERSION and EXPECTED_COUNTS together in one reviewed edit.
DATASET_VERSION = "mvp_v1_3"
EXPECTED_COUNTS = {"menus": 227, "hours": 10, "holidays": 11, "metadata": 11}

HEADER = (
    "-- Aggrey MVP seed data. Generated from Aggrey_MVP_Data.xlsx without changing\n"
    "-- any menu item, serving window, holiday note or metadata value.\n"
    "-- Run after schema.sql, against fresh empty tables.\n"
    "-- Re-running is duplicate-safe, but does NOT remove deleted/renamed menu items.\n"
    "-- This is an INITIAL IMPORT, not an ongoing synchronization process."
)


def read_sheet(wb: "openpyxl.Workbook", name: str) -> list[dict[str, object]]:
    """Return the sheet's data rows as dicts keyed by the expected column names."""
    if name not in wb.sheetnames:
        sys.exit(f"ERROR: workbook is missing the '{name}' sheet")
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    header = [h.strip() if isinstance(h, str) else h for h in rows[0]]
    if header != COLUMNS[name]:
        sys.exit(f"ERROR: '{name}' sheet columns {header} != expected {COLUMNS[name]}")
    records: list[dict[str, object]] = []
    for raw in rows[1:]:
        # Skip fully blank trailing rows that spreadsheets sometimes keep.
        if all(cell is None or (isinstance(cell, str) and cell.strip() == "") for cell in raw):
            continue
        records.append(dict(zip(COLUMNS[name], raw, strict=True)))
    return records


def q(value: object) -> str:
    """Render a value as a single-quoted SQL string, or NULL. Escapes quotes."""
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def num_or_null(value: object) -> str:
    """Render an integer value unquoted, or NULL (used for numeric columns)."""
    if value is None:
        return "NULL"
    return str(int(value))


def menu_row(r: dict[str, object]) -> str:
    return (
        f"    ({q(r['menu_scope'])}, {num_or_null(r['day_of_week'])}, "
        f"{q(r['meal_type'])}, {q(r['station'])}, {q(r['food_item'])}, "
        f"{num_or_null(r['display_order'])})"
    )


def hours_row(r: dict[str, object]) -> str:
    return (
        f"    ({q(r['schedule_type'])}, {q(r['meal_type'])}, "
        f"{q(r['start_time'])}, {q(r['end_time'])})"
    )


def holiday_row(r: dict[str, object]) -> str:
    return (
        f"    ({q(r['name'])}, {q(r['start_date'])}, {q(r['end_date'])}, "
        f"{q(r['schedule_type'])}, {q(r['verification_status'])}, {q(r['notes'])})"
    )


def metadata_row(r: dict[str, object]) -> str:
    return f"    ({q(r['key'])}, {q(r['value'])}, {q(r['notes'])})"


def insert_block(header_line: str, rows: list[str], conflict: str) -> str:
    """One INSERT statement: header line, comma-joined value rows, conflict clause."""
    return header_line + "\n" + ",\n".join(rows) + "\n" + conflict


def render(
    menus: list[dict[str, object]],
    hours: list[dict[str, object]],
    holidays: list[dict[str, object]],
    metadata: list[dict[str, object]],
) -> str:
    menu_cols = "INSERT INTO public.menus (menu_scope, day_of_week, meal_type, station, food_item, display_order) VALUES"
    menu_blocks = [
        insert_block(menu_cols, [menu_row(r) for r in menus[i : i + MENU_BATCH_SIZE]],
                     "ON CONFLICT DO NOTHING;")
        for i in range(0, len(menus), MENU_BATCH_SIZE)
    ]
    menu_section = "\n\n".join(menu_blocks)

    hours_block = insert_block(
        "INSERT INTO public.hours (schedule_type, meal_type, start_time, end_time) VALUES",
        [hours_row(r) for r in hours],
        "ON CONFLICT (schedule_type, meal_type) DO UPDATE SET\n"
        "    start_time = EXCLUDED.start_time,\n"
        "    end_time = EXCLUDED.end_time;",
    )
    holidays_block = insert_block(
        "INSERT INTO public.holidays (name, start_date, end_date, schedule_type, verification_status, notes) VALUES",
        [holiday_row(r) for r in holidays],
        "ON CONFLICT (name, start_date, end_date) DO UPDATE SET\n"
        "    schedule_type = EXCLUDED.schedule_type,\n"
        "    verification_status = EXCLUDED.verification_status,\n"
        "    notes = EXCLUDED.notes;",
    )
    metadata_block = insert_block(
        'INSERT INTO public.metadata ("key", value, notes) VALUES',
        [metadata_row(r) for r in metadata],
        'ON CONFLICT ("key") DO UPDATE SET\n'
        "    value = EXCLUDED.value,\n"
        "    notes = EXCLUDED.notes;",
    )

    return (
        HEADER
        + "\n\nBEGIN;\n\n"
        + menu_section
        + "\n" + hours_block
        + "\n" + holidays_block
        + "\n" + metadata_block
        + "\nCOMMIT;\n"
    )


def _blank(value: object) -> bool:
    """True if the value is missing or an empty/whitespace-only string."""
    return value is None or (isinstance(value, str) and value.strip() == "")


def _minutes(value: object) -> int | None:
    """Parse 'HH:MM' or 'HH:MM:SS' text to minutes, or None if malformed.

    Requires a text cell with valid ranges (0<=h<24, 0<=m<60, 0<=s<60). Native
    Excel time cells (datetime.time) and junk like '25:99' or '17:00:x' return None,
    so the seed never silently renders '07:30' as '07:30:00'.
    """
    if not isinstance(value, str):
        return None
    parts = value.strip().split(":")
    if len(parts) not in (2, 3) or not all(p.isdigit() for p in parts):
        return None
    h, m, *rest = (int(p) for p in parts)
    s = rest[0] if rest else 0
    if not (0 <= h < 24 and 0 <= m < 60 and 0 <= s < 60):
        return None
    return h * 60 + m


def _date(value: object) -> date | None:
    """Parse a whole 'YYYY-MM-DD' text value, or None if malformed.

    Parses the entire stripped string (no truncation), so '2026-09-07x' is rejected,
    and requires text so native Excel date cells (datetime) are rejected too.
    """
    if not isinstance(value, str):
        return None
    try:
        return date.fromisoformat(value.strip())
    except ValueError:
        return None


def validate(
    menus: list[dict[str, object]],
    hours: list[dict[str, object]],
    holidays: list[dict[str, object]],
    metadata: list[dict[str, object]],
    expected_counts: dict[str, int] = EXPECTED_COUNTS,
) -> list[str]:
    """Return every problem found in the workbook. Empty list means it is clean."""
    errors: list[str] = []

    # Row counts must match the pinned dataset version.
    for sheet, rows in (("menus", menus), ("hours", hours),
                        ("holidays", holidays), ("metadata", metadata)):
        if len(rows) != expected_counts[sheet]:
            errors.append(
                f"{sheet}: {len(rows)} rows; expected {expected_counts[sheet]} for "
                f"{DATASET_VERSION} (bump DATASET_VERSION and EXPECTED_COUNTS if the "
                f"data changed on purpose)"
            )

    # --- menus ---
    seen_menu: dict[tuple, int] = {}
    for i, r in enumerate(menus, start=1):
        scope, dow, meal = r["menu_scope"], r["day_of_week"], r["meal_type"]
        if scope not in MENU_SCOPES:
            errors.append(f"menus row {i}: menu_scope '{scope}' not in {sorted(MENU_SCOPES)}")
        if meal not in MEAL_TYPES:
            errors.append(f"menus row {i}: meal_type '{meal}' not in {sorted(MEAL_TYPES)}")
        if _blank(r["station"]):
            errors.append(f"menus row {i}: station is blank")
        if _blank(r["food_item"]):
            errors.append(f"menus row {i}: food_item is blank")
        order = r["display_order"]
        if not isinstance(order, int) or isinstance(order, bool) or order <= 0:
            errors.append(f"menus row {i}: display_order must be a positive integer, got {r['display_order']!r}")

        if scope == "holiday_sample":
            if dow is not None:
                errors.append(f"menus row {i}: holiday_sample row must have a blank day_of_week")
            if meal not in {"brunch", "dinner"}:
                errors.append(f"menus row {i}: holiday_sample meal_type must be brunch or dinner")
        elif scope == "weekly":
            if not isinstance(dow, int) or isinstance(dow, bool) or dow < 1 or dow > 7:
                errors.append(f"menus row {i}: weekly day_of_week must be 1-7, got {dow!r}")
            elif dow <= 5 and meal not in {"breakfast", "lunch", "dinner"}:
                errors.append(f"menus row {i}: weekday {dow} meal_type must be breakfast/lunch/dinner")
            elif dow >= 6 and meal not in {"brunch", "dinner"}:
                errors.append(f"menus row {i}: weekend day {dow} meal_type must be brunch/dinner")

        key = (scope, dow, meal, str(r["station"]).strip().lower(), str(r["food_item"]).strip().lower())
        if key in seen_menu:
            errors.append(f"menus row {i}: duplicate of row {seen_menu[key]} (same day/meal/station/item)")
        else:
            seen_menu[key] = i

    # --- hours ---
    seen_hours: dict[tuple, int] = {}
    for i, r in enumerate(hours, start=1):
        sched, meal = r["schedule_type"], r["meal_type"]
        if sched not in SCHEDULE_TYPES:
            errors.append(f"hours row {i}: schedule_type '{sched}' not in {sorted(SCHEDULE_TYPES)}")
        if meal not in MEAL_TYPES:
            errors.append(f"hours row {i}: meal_type '{meal}' not in {sorted(MEAL_TYPES)}")
        if sched == "weekday" and meal not in {"breakfast", "lunch", "lite_lunch", "dinner"}:
            errors.append(f"hours row {i}: weekday meal_type must be breakfast/lunch/lite_lunch/dinner")
        if sched in {"weekend", "holiday", "break"} and meal not in {"brunch", "dinner"}:
            errors.append(f"hours row {i}: {sched} meal_type must be brunch/dinner")
        start, end = _minutes(r["start_time"]), _minutes(r["end_time"])
        if start is None or end is None:
            errors.append(f"hours row {i}: start/end time must be HH:MM, got {r['start_time']!r}/{r['end_time']!r}")
        elif end <= start:
            errors.append(f"hours row {i}: end_time {r['end_time']} is not after start_time {r['start_time']}")
        if (sched, meal) in seen_hours:
            errors.append(f"hours row {i}: duplicate schedule/meal of row {seen_hours[(sched, meal)]}")
        else:
            seen_hours[(sched, meal)] = i

    # --- holidays ---
    seen_hol: dict[tuple, int] = {}
    ranges: list[tuple[int, str, date, date]] = []
    for i, r in enumerate(holidays, start=1):
        if _blank(r["name"]):
            errors.append(f"holidays row {i}: name is blank")
        if _blank(r["verification_status"]):
            errors.append(f"holidays row {i}: verification_status is blank")
        if r["schedule_type"] not in {"holiday", "break"}:
            errors.append(f"holidays row {i}: schedule_type must be holiday or break, got {r['schedule_type']!r}")
        start, end = _date(r["start_date"]), _date(r["end_date"])
        if start is None or end is None:
            errors.append(f"holidays row {i}: dates must be YYYY-MM-DD, got {r['start_date']!r}/{r['end_date']!r}")
        elif end < start:
            errors.append(f"holidays row {i}: end_date {r['end_date']} is before start_date {r['start_date']}")
        else:
            ranges.append((i, str(r["name"]), start, end))
        dup_key = (r["name"], r["start_date"], r["end_date"])
        if dup_key in seen_hol:
            errors.append(f"holidays row {i}: duplicate name/dates of row {seen_hol[dup_key]}")
        else:
            seen_hol[dup_key] = i
    # Overlapping ranges (inclusive), every pair.
    for a in range(len(ranges)):
        for b in range(a + 1, len(ranges)):
            ia, na, sa, ea = ranges[a]
            ib, nb, sb, eb = ranges[b]
            if sa <= eb and sb <= ea:
                errors.append(f"holidays: '{na}' (row {ia}) overlaps '{nb}' (row {ib})")

    # --- metadata ---
    seen_keys: dict[str, int] = {}
    for i, r in enumerate(metadata, start=1):
        if _blank(r["value"]):
            errors.append(f"metadata row {i}: value is blank")
        if _blank(r["key"]):
            errors.append(f"metadata row {i}: key is blank")
            continue  # don't also report a confusing duplicate 'None' key
        k = str(r["key"]).strip()
        if k in seen_keys:
            errors.append(f"metadata row {i}: duplicate key '{k}' (also row {seen_keys[k]})")
        else:
            seen_keys[k] = i
    version = next(
        (str(r["value"]).strip() for r in metadata if str(r["key"]).strip() == "dataset_version"),
        None,
    )
    if version != DATASET_VERSION:
        errors.append(f"metadata: dataset_version is {version!r}; this generator expects {DATASET_VERSION!r}")

    return errors


def mirror_cell(value: object) -> str:
    """Render a workbook value the way the Markdown mirror shows it (blank for None)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_mirror(md_text: str) -> dict[str, list[tuple[str, ...]]]:
    """Return the Markdown mirror's data rows per sheet (as tuples of cell strings)."""
    tables: dict[str, list[tuple[str, ...]]] = {sheet: [] for sheet in COLUMNS}
    current: str | None = None
    for line in md_text.splitlines():
        stripped = line.strip()
        if stripped.startswith("## "):
            current = SECTION_TO_SHEET.get(stripped[3:].strip())
            continue
        if current is None or not stripped.startswith("|"):
            continue
        cells = [c.strip() for c in stripped.strip("|").split("|")]
        if all(set(c) <= {"-", ":"} and c for c in cells):  # separator row | --- |
            continue
        if cells == COLUMNS[current]:  # the column-header row
            continue
        tables[current].append(tuple(cells))
    return tables


def check_mirror(
    menus: list[dict[str, object]],
    hours: list[dict[str, object]],
    holidays: list[dict[str, object]],
    metadata: list[dict[str, object]],
) -> list[str]:
    """Return differences between the workbook and the Markdown mirror (empty = match)."""
    if not MIRROR_FILE.exists():
        return [f"mirror file not found at {MIRROR_FILE}"]
    md_tables = parse_mirror(MIRROR_FILE.read_text(encoding="utf-8"))
    sheets = {"menus": menus, "hours": hours, "holidays": holidays, "metadata": metadata}

    errors: list[str] = []
    for name, rows in sheets.items():
        wb_rows = [tuple(mirror_cell(r[col]) for col in COLUMNS[name]) for r in rows]
        md_rows = md_tables[name]
        # Flag any mirror row with the wrong number of columns.
        for bad in (c for c in md_rows if len(c) != len(COLUMNS[name])):
            errors.append(f"{name}: mirror row has {len(bad)} columns, expected {len(COLUMNS[name])}: {bad}")
        wb_count, md_count = Counter(wb_rows), Counter(md_rows)
        for row in wb_count - md_count:
            errors.append(f"{name}: in workbook but not in mirror: {row}")
        for row in md_count - wb_count:
            errors.append(f"{name}: in mirror but not in workbook: {row}")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate db/seed.sql from the workbook.")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--check", action="store_true",
                      help="exit non-zero if the committed seed.sql would change")
    mode.add_argument("--check-mirror", action="store_true",
                      help="exit non-zero if data/Aggrey_MVP_Data.md differs from the workbook")
    args = parser.parse_args()

    if not WORKBOOK.exists():
        sys.exit(f"ERROR: workbook not found at {WORKBOOK}")

    import openpyxl  # imported here so tests can use this module without the dependency

    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    menus = read_sheet(wb, "menus")
    hours = read_sheet(wb, "hours")
    holidays = read_sheet(wb, "holidays")
    metadata = read_sheet(wb, "metadata")

    errors = validate(menus, hours, holidays, metadata)
    if errors:
        print(f"ERROR: workbook failed validation ({len(errors)} problem(s)):", file=sys.stderr)
        for e in errors:
            print(f"  - {e}", file=sys.stderr)
        return 1

    if args.check_mirror:
        diffs = check_mirror(menus, hours, holidays, metadata)
        if not diffs:
            print("OK: data/Aggrey_MVP_Data.md matches the workbook.")
            return 0
        print(f"ERROR: Markdown mirror is out of sync with the workbook ({len(diffs)} difference(s)):")
        for d in diffs:
            print(f"  - {d}")
        return 1

    generated = render(menus, hours, holidays, metadata)

    if args.check:
        current = SEED_FILE.read_text(encoding="utf-8") if SEED_FILE.exists() else ""
        if generated == current:
            print("OK: db/seed.sql matches the workbook.")
            return 0
        diff = difflib.unified_diff(
            current.splitlines(), generated.splitlines(),
            fromfile="db/seed.sql (committed)", tofile="db/seed.sql (from workbook)",
            lineterm="",
        )
        print("ERROR: db/seed.sql is out of date with the workbook. Regenerate it.\n")
        print("\n".join(diff))
        return 1

    SEED_FILE.write_text(generated, encoding="utf-8")
    print(f"Wrote {SEED_FILE.relative_to(REPO_ROOT)} from the workbook.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
